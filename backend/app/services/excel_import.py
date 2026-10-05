"""Import tolérant de Suivi_Licence.xlsx avec UPSERT et rapport d'anomalies."""

import io
import json
from dataclasses import dataclass, field
from datetime import date

from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.ext.asyncio import AsyncSession

from app.models import Asset, Contract, ImportLog, LicenseAssignment, Vendor
from app.services import excel_layout as L
from app.services.parsing import PARSERS, ParseError, parse_str
from app.services.references import next_reference, sync_sequence

HEADER_SCAN_ROWS = 15
VENDOR_FIELDS = ("contact_person", "email_support", "phone", "website")
MAX_LEN = {"reference": 100, "name": 255}


class ImportFileError(ValueError):
    pass


@dataclass
class SheetStats:
    sheet: str
    matched_as: str | None
    rows_analyzed: int = 0
    imported: int = 0
    created: int = 0
    updated: int = 0
    skipped: int = 0
    note: str | None = None


@dataclass
class Report:
    filename: str
    sheets: list[SheetStats] = field(default_factory=list)
    warnings: list[dict] = field(default_factory=list)
    errors: list[dict] = field(default_factory=list)
    missing_sheets: list[str] = field(default_factory=list)
    unknown_sheets: list[str] = field(default_factory=list)
    vendors_created: int = 0

    def warn(self, sheet, row, msg):
        self.warnings.append({"sheet": sheet, "row": row, "message": msg})

    def error(self, sheet, row, msg):
        self.errors.append({"sheet": sheet, "row": row, "message": msg})

    def as_dict(self) -> dict:
        return {
            "filename": self.filename,
            "rows_analyzed": sum(s.rows_analyzed for s in self.sheets),
            "rows_imported": sum(s.imported for s in self.sheets),
            "rows_skipped": sum(s.skipped for s in self.sheets),
            "created": sum(s.created for s in self.sheets),
            "updated": sum(s.updated for s in self.sheets),
            "vendors_created": self.vendors_created,
            "warnings_count": len(self.warnings),
            "errors_count": len(self.errors),
            "sheets": [s.__dict__ for s in self.sheets],
            "missing_sheets": self.missing_sheets,
            "unknown_sheets": self.unknown_sheets,
            "warnings": self.warnings,
            "errors": self.errors,
        }


def detect_header(rows: list[tuple], sheet: L.Sheet) -> tuple[int, dict[int, L.Col], list[str]]:
    """Retourne (index ligne d'en-tête, {index colonne: Col}, colonnes non reconnues)."""
    best = (-1, {}, [])
    for idx, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        mapping: dict[int, L.Col] = {}
        used: set[str] = set()
        cells = [(i, L.norm(v)) for i, v in enumerate(row) if v is not None and L.norm(v)]
        # 1er passage : libellé exact, 2e passage : alias
        for exact in (True, False):
            for i, n in cells:
                if i in mapping:
                    continue
                for col in sheet.columns:
                    if col.key in used:
                        continue
                    if (exact and n == L.norm(col.label)) or (not exact and n in col.aliases):
                        mapping[i] = col
                        used.add(col.key)
                        break
        if len(mapping) > len(best[1]):
            unknown = [str(row[i]) for i, _ in cells if i not in mapping]
            best = (idx, mapping, unknown)
    return best


class Importer:
    def __init__(self, db: AsyncSession, report: Report):
        self.db = db
        self.report = report
        self.vendors: dict[str, Vendor] = {}
        self.contracts: dict[str, Contract] = {}
        # Références explicites présentes dans le fichier : jamais attribuées à une ligne sans référence
        self.reserved: dict[str, set[str]] = {}

    async def load_caches(self):
        for v in (await self.db.execute(select(Vendor))).scalars():
            self.vendors[L.norm(v.name)] = v
        for c in (await self.db.execute(select(Contract))).scalars():
            self.contracts[c.reference.strip().upper()] = c

    async def vendor(self, name: str | None) -> Vendor | None:
        if not name:
            return None
        key = L.norm(name)
        if not key:
            return None
        v = self.vendors.get(key)
        if v is None:
            v = Vendor(name=name[:255], reference=await next_reference(self.db, "VENDOR"))
            self.db.add(v)
            await self.db.flush()
            self.vendors[key] = v
            self.report.vendors_created += 1
        return v

    async def contract(self, ref: str | None, vendor: Vendor | None, sheet: str, row: int) -> Contract | None:
        if not ref:
            return None
        key = ref.strip().upper()
        c = self.contracts.get(key)
        if c is None:
            c = Contract(reference=ref[:100], vendor_id=vendor.id if vendor else None, status="À compléter")
            self.db.add(c)
            await self.db.flush()
            self.contracts[key] = c
            self.report.warn(sheet, row, f"contrat « {ref} » absent de l'onglet contrats : créé avec des informations minimales")
        return c

    def parse_row(self, row: tuple, mapping: dict[int, L.Col], sheet: str, rownum: int) -> dict:
        data: dict = {}
        for i, col in mapping.items():
            if not col.importable or i >= len(row):
                continue
            try:
                data[col.key] = PARSERS[col.kind](row[i])
            except ParseError as exc:
                data[col.key] = None
                self.report.warn(sheet, rownum, f"{col.label} : {exc} → valeur ignorée")
        return data

    async def import_sheet(self, ws, ws_formulas, sheet: L.Sheet):
        stats = SheetStats(sheet=ws.title, matched_as=sheet.name)
        self.report.sheets.append(stats)
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            stats.note = "onglet vide"
            return
        header_idx, mapping, unknown = detect_header(rows, sheet)
        keys = {c.key for c in mapping.values()}
        if header_idx < 0 or not ({"reference", "name"} & keys):
            stats.note = "en-têtes non reconnus"
            self.report.error(
                ws.title, None, "Impossible d'identifier la ligne d'en-têtes (colonnes Référence / Nom introuvables)"
            )
            return
        for u in unknown:
            self.report.warn(ws.title, header_idx + 1, f"colonne non reconnue « {u} » ignorée")
        formula_rows = list(ws_formulas.iter_rows(values_only=True)) if ws_formulas is not None else []

        ref_idx = next((i for i, c in mapping.items() if c.key == "reference"), None)
        if ref_idx is not None:
            reserved = self.reserved.setdefault(sheet.category, set())
            for r in rows[header_idx + 1 :]:
                if ref_idx < len(r) and (v := parse_str(r[ref_idx])):
                    reserved.add(v.strip().upper())

        seen: dict[str, int] = {}
        for offset, row in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
            if all(v is None or (isinstance(v, str) and not v.strip()) for v in row):
                continue
            # Formules sans valeur en cache : signaler
            if formula_rows and offset - 1 < len(formula_rows):
                frow = formula_rows[offset - 1]
                for i, col in mapping.items():
                    if (
                        col.importable
                        and i < len(row)
                        and row[i] is None
                        and i < len(frow)
                        and isinstance(frow[i], str)
                        and frow[i].startswith("=")
                    ):
                        self.report.warn(
                            ws.title, offset, f"{col.label} : formule sans valeur calculée ({frow[i][:40]}) → ignorée"
                        )
            stats.rows_analyzed += 1
            data = self.parse_row(row, mapping, ws.title, offset)
            if all(v is None for v in data.values()):
                stats.rows_analyzed -= 1  # ligne ne contenant que des colonnes calculées
                continue
            ref, name = data.get("reference"), data.get("name")
            generated = False
            if sheet.category != "CONTRACT":
                if not ref and not name:
                    stats.skipped += 1
                    self.report.error(ws.title, offset, "ligne incomplète : ni référence ni nom → ignorée")
                    continue
                if not name:
                    data["name"] = name = ref
                    self.report.warn(ws.title, offset, "nom manquant : la référence est utilisée comme nom")
                generated = not ref
                identity = ref or f"nom:{L.norm(name)}"
            else:
                # Contrat sans référence : identifiable seulement s'il porte un fournisseur, un périmètre ou un type.
                if not ref and not (data.get("vendor") or data.get("scope") or data.get("type")):
                    stats.skipped += 1
                    self.report.error(ws.title, offset, "référence contrat manquante et ligne non identifiable → ligne ignorée")
                    continue
                generated = not ref
                identity = ref or f"ctr:{L.norm(data.get('vendor'))}:{L.norm(data.get('scope') or data.get('type'))}"
            data["reference"] = ref  # None si à générer au moment de l'écriture
            for k, n in MAX_LEN.items():
                if data.get(k) and len(data[k]) > n:
                    data[k] = data[k][:n]
                    self.report.warn(ws.title, offset, f"{k} tronqué à {n} caractères")
            dkey = identity.strip().upper()
            if dkey in seen:
                stats.skipped += 1
                what = f"du nom « {name} »" if generated and sheet.category != "CONTRACT" else f"de la référence « {ref} »"
                self.report.warn(ws.title, offset, f"doublon {what} (déjà présent ligne {seen[dkey]}) → ligne ignorée")
                continue
            seen[dkey] = offset
            try:
                async with self.db.begin_nested():
                    if sheet.category == "CONTRACT":
                        created = await self.upsert_contract(data, ws.title, offset)
                    else:
                        created = await self.upsert_asset(sheet.category, data, ws.title, offset)
                    if generated:
                        self.report.warn(
                            ws.title,
                            offset,
                            f"référence manquante : « {data['reference']} » {'générée' if created else 'déjà attribuée (ligne reconnue)'}",
                        )
            except SQLAlchemyError as exc:
                stats.skipped += 1
                self.report.error(
                    ws.title, offset, f"erreur base de données : {str(exc.orig if hasattr(exc, 'orig') else exc)[:200]}"
                )
                continue
            stats.imported += 1
            stats.created += int(created)
            stats.updated += int(not created)

    async def upsert_contract(self, data: dict, sheet: str, row: int) -> bool:
        vendor = await self.vendor(data.pop("vendor", None))
        if vendor:
            for f in VENDOR_FIELDS:
                if data.get(f) and not getattr(vendor, f):
                    setattr(vendor, f, data[f][:255])
        for f in VENDOR_FIELDS:
            data.pop(f, None)
        if vendor is None and all(v is None for k, v in data.items() if k != "reference"):
            self.report.warn(sheet, row, "ligne incomplète : seule la référence du contrat est renseignée")
        if not data["reference"]:
            # Ré-import d'une ligne sans référence : on retrouve le contrat par fournisseur + périmètre/type
            label = (data.get("scope") or data.get("type") or "").lower()
            for existing in self.contracts.values():
                same_vendor = (existing.vendor_id or None) == (vendor.id if vendor else None)
                if same_vendor and (existing.scope or existing.type or "").lower() == label:
                    data["reference"] = existing.reference
                    break
            else:
                data["reference"] = await next_reference(self.db, "CONTRACT", self.reserved.get("CONTRACT"))
        key = data["reference"].strip().upper()
        c = self.contracts.get(key)
        created = c is None
        if created:
            c = Contract(reference=data["reference"])
            self.db.add(c)
        if vendor:
            c.vendor_id = vendor.id
        for k, v in data.items():
            if v is not None and k != "reference":
                setattr(c, k, v)
        if c.start_date and c.end_date and c.end_date < c.start_date:
            self.report.warn(sheet, row, "date de fin antérieure à la date de début")
        await self.db.flush()
        self.contracts[key] = c
        return created

    async def upsert_asset(self, category: str, data: dict, sheet: str, row: int) -> bool:
        vendor = await self.vendor(data.pop("vendor", None))
        contract = await self.contract(data.pop("contract", None), vendor, sheet, row)
        used = data.pop("used_quantity", None)
        if data["reference"]:
            lookup = select(Asset).where(Asset.category == category, Asset.reference == data["reference"])
        else:
            # Référence absente : on retrouve l'actif par son nom (ré-import idempotent), sinon on en génère une
            lookup = select(Asset).where(Asset.category == category, func.lower(Asset.name) == data["name"].lower()).limit(1)
        existing = (await self.db.execute(lookup)).scalar_one_or_none()
        if existing is not None:
            data["reference"] = existing.reference
        elif not data["reference"]:
            data["reference"] = await next_reference(self.db, category, self.reserved.get(category))
        created = existing is None
        asset = existing or Asset(category=category, reference=data["reference"])
        if created:
            self.db.add(asset)
        if vendor:
            asset.vendor_id = vendor.id
        if contract:
            asset.contract_id = contract.id
        for k, v in data.items():
            if v is not None and k != "reference":
                setattr(asset, k, v)
        if asset.start_date and asset.end_date and asset.end_date < asset.start_date:
            self.report.warn(sheet, row, "date d'expiration antérieure à la date de début")
        await self.db.flush()

        if category == "LICENCE" and used:
            if asset.total_quantity is None:
                asset.total_quantity = used
                self.report.warn(sheet, row, "quantité totale manquante : alignée sur la quantité utilisée")
            has_assign = (
                await self.db.execute(select(LicenseAssignment.id).where(LicenseAssignment.asset_id == asset.id).limit(1))
            ).first()
            if not has_assign:
                qty = used
                if used > asset.total_quantity:
                    qty = asset.total_quantity
                    self.report.warn(
                        sheet,
                        row,
                        f"sur-allocation dans Excel ({used} utilisées > {asset.total_quantity}) : affectation plafonnée à {qty}",
                    )
                if qty > 0:
                    self.db.add(
                        LicenseAssignment(
                            asset_id=asset.id,
                            quantity=qty,
                            assigned_to_department=asset.user_department,
                            assigned_date=date.today(),
                            notes="Affectation globale reprise de l'import Excel",
                        )
                    )
                    await self.db.flush()
        return created


async def import_workbook(
    db: AsyncSession, content: bytes, filename: str, user: str | None = None, file_size: int | None = None
) -> dict:
    report = Report(filename=filename)
    try:
        wb = load_workbook(io.BytesIO(content), data_only=True, read_only=False)
        wbf = load_workbook(io.BytesIO(content), data_only=False, read_only=False)
    except Exception as exc:  # noqa: BLE001 - tout fichier illisible
        raise ImportFileError(f"Fichier Excel illisible : {exc}") from exc

    matched: dict[str, tuple] = {}
    for ws in wb.worksheets:
        sheet = L.match_sheet(ws.title)
        if sheet is None:
            report.unknown_sheets.append(ws.title)
            report.warn(ws.title, None, "onglet non reconnu : ignoré")
        elif sheet.name not in matched:
            matched[sheet.name] = (ws, wbf[ws.title], sheet)
    for sh in L.ALL_SHEETS:
        if sh.name not in matched:
            report.missing_sheets.append(sh.name)
            report.warn(sh.name, None, "onglet manquant dans le fichier")
    if not any(sh.importable for _, _, sh in matched.values()):
        raise ImportFileError(
            "Aucun onglet métier reconnu (Licences_Logicielles, Certificats, Materiels, Applications, Contrats_Fournisseurs)"
        )

    imp = Importer(db, report)
    await imp.load_caches()
    order = [L.CONTRATS, *L.ASSET_SHEETS]  # contrats d'abord pour lier les actifs
    for sh in order:
        if sh.name in matched:
            ws, wsf, _ = matched[sh.name]
            await imp.import_sheet(ws, wsf, sh)
    for sh in (L.PLANNING, L.DASHBOARD):
        if sh.name in matched:
            report.sheets.append(
                SheetStats(
                    sheet=matched[sh.name][0].title,
                    matched_as=sh.name,
                    note="onglet calculé : régénéré depuis la base, non importé",
                )
            )

    for kind in ("LICENCE", "CERTIFICAT", "MATERIEL", "APPLICATION", "CONTRACT", "VENDOR"):
        await sync_sequence(db, kind)

    result = report.as_dict()
    db.add(
        ImportLog(
            filename=filename[:255],
            rows_analyzed=result["rows_analyzed"],
            rows_imported=result["rows_imported"],
            rows_skipped=result["rows_skipped"],
            warnings_count=result["warnings_count"],
            errors_count=result["errors_count"],
            report=json.dumps(result, default=str, ensure_ascii=False),
            imported_by=user,
            file_size=file_size,
            status="ERRORS" if result["errors_count"] else ("WARNINGS" if result["warnings_count"] else "SUCCESS"),
        )
    )
    await db.commit()
    return result
