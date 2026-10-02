"""Export Excel reproduisant les onglets métier de Suivi_Licence.xlsx."""

import io
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.ext.asyncio import AsyncSession

from app.core import lifecycle
from app.core.config import settings
from app.services import excel_layout as L
from app.services.dashboard import build_dashboard, filter_renewals, load_all, renewal_items

HEADER_FILL = PatternFill("solid", fgColor="1E3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14, color="1E3A5F")
THIN = Side(style="thin", color="D0D7E2")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
STATUS_FILLS = {
    "EXPIRÉ": "F8D7DA",
    "CRITIQUE": "FDE2C8",
    "ALERTE": "FFF3CD",
    "OK": "D4EDDA",
}
DATE_FMT = "DD/MM/YYYY"
MONEY_FMT = f'#,##0 "{settings.currency_label}"'


def _asset_value(asset, col: L.Col):
    k = col.key
    if k == "vendor":
        return asset.vendor_name
    if k == "contract":
        return asset.contract_reference
    if k == "used_quantity":
        return asset.assigned_quantity
    if k == "status":
        return lifecycle.STATUS_LABELS[asset.status]
    if k == "usage_rate":
        return asset.usage_rate
    v = getattr(asset, k, None)
    return float(v) if col.kind == "money" and v is not None else v


def _contract_value(c, col: L.Col):
    k = col.key
    if k == "vendor":
        return c.vendor_name
    if k in ("contact_person", "email_support", "phone", "website"):
        return getattr(c.vendor, k) if c.vendor else None
    if k == "lifecycle_status":
        return lifecycle.STATUS_LABELS[c.lifecycle_status]
    v = getattr(c, k, None)
    return float(v) if col.kind == "money" and v is not None else v


def _write_table(ws, columns: list[L.Col], rows: list[list], start_row: int = 1):
    for j, col in enumerate(columns, start=1):
        cell = ws.cell(row=start_row, column=j, value=col.label)
        cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, values in enumerate(rows, start=start_row + 1):
        for j, (col, v) in enumerate(zip(columns, values, strict=False), start=1):
            cell = ws.cell(row=i, column=j, value=v)
            cell.border = BORDER
            if isinstance(v, (date, datetime)):
                cell.number_format = DATE_FMT
            elif col.kind == "money" or col.key == "budget":
                cell.number_format = MONEY_FMT
    for j, col in enumerate(columns, start=1):
        width = max([len(col.label)] + [len(str(r[j - 1])) for r in rows[:200] if r[j - 1] is not None])
        ws.column_dimensions[get_column_letter(j)].width = min(max(width + 2, 10), 45)
    ws.row_dimensions[start_row].height = 30
    ws.freeze_panes = ws.cell(row=start_row + 1, column=3)
    if rows:
        ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(columns))}{start_row + len(rows)}"
        status_cols = [j for j, c in enumerate(columns, 1) if c.key in ("status", "lifecycle_status")]
        for j in status_cols:
            letter = get_column_letter(j)
            rng = f"{letter}{start_row + 1}:{letter}{start_row + len(rows)}"
            for label, color in STATUS_FILLS.items():
                ws.conditional_formatting.add(
                    rng, CellIsRule(operator="equal", formula=[f'"{label}"'], fill=PatternFill("solid", fgColor=color))
                )


def _write_dashboard(ws, d: dict):
    ws["A1"] = "Tableau de bord — Suivi des licences, certificats, matériels, applications et contrats"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Généré le {lifecycle.today().strftime('%d/%m/%Y')} depuis la base de données"
    ws["A2"].font = Font(italic=True, color="666666")
    k, dl, li, fi = d["kpis"], d["deadlines"], d["licenses"], d["finance"]
    blocks = [
        (
            "Inventaire",
            [
                ("Total actifs", k["total_assets"]),
                ("Licences", k["licences"]),
                ("Certificats", k["certificats"]),
                ("Matériels", k["materiels"]),
                ("Applications", k["applications"]),
                ("Contrats", k["contracts"]),
                ("Fournisseurs", k["vendors"]),
            ],
        ),
        (
            "Échéances (actifs)",
            [
                ("Expirés", dl["expired"]),
                ("Critiques (≤ 30 j)", dl["critical"]),
                ("À surveiller (≤ 90 j)", dl["warning"]),
                ("OK (> 90 j)", dl["ok"]),
                ("Sans date", dl["unknown"]),
            ],
        ),
        (
            "Licences",
            [
                ("Total licences", li["total"]),
                ("Licences utilisées", li["used"]),
                ("Licences disponibles", li["available"]),
                ("Taux d'utilisation (%)", li["usage_rate"]),
            ],
        ),
        (
            "Finances",
            [
                ("Coût annuel total actifs", fi["annual_cost_total"]),
                ("Montant annuel contrats", fi["contracts_annual_total"]),
                ("Budget estimé", fi["budget_estimated_total"]),
                ("Budget renouvellement 12 mois", fi["renewal_budget_12m"]),
            ],
        ),
    ]
    row = 4
    for title, kv in blocks:
        c = ws.cell(row=row, column=1, value=title)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        ws.cell(row=row, column=2).fill = HEADER_FILL
        row += 1
        for label, value in kv:
            ws.cell(row=row, column=1, value=label).border = BORDER
            vc = ws.cell(row=row, column=2, value=value)
            vc.border = BORDER
            if label.startswith(("Coût", "Montant", "Budget")):
                vc.number_format = MONEY_FMT
            row += 1
        row += 1
    ws.cell(row=row, column=1, value="Répartition des statuts par catégorie").font = Font(bold=True, color="1E3A5F")
    row += 1
    headers = ["Catégorie"] + [lifecycle.STATUS_LABELS[s] for s in lifecycle.STATUSES]
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=j, value=h)
        c.fill, c.font, c.border = HEADER_FILL, HEADER_FONT, BORDER
    for r in d["status_by_category"]:
        row += 1
        ws.cell(row=row, column=1, value=r["label"]).border = BORDER
        for j, s in enumerate(lifecycle.STATUSES, 2):
            ws.cell(row=row, column=j, value=r[s]).border = BORDER
    ws.column_dimensions["A"].width = 40
    for col in "BCDEF":
        ws.column_dimensions[col].width = 16


async def export_workbook(db: AsyncSession) -> bytes:
    assets, contracts = await load_all(db)
    wb = Workbook()
    wb.remove(wb.active)
    for sheet in L.ASSET_SHEETS:
        ws = wb.create_sheet(sheet.name)
        rows = [
            [_asset_value(a, c) for c in sheet.columns]
            for a in sorted((a for a in assets if a.category == sheet.category), key=lambda a: a.reference)
        ]
        _write_table(ws, sheet.columns, rows)
    ws = wb.create_sheet(L.CONTRATS.name)
    _write_table(
        ws,
        L.CONTRATS.columns,
        [[_contract_value(c, col) for col in L.CONTRATS.columns] for c in sorted(contracts, key=lambda c: c.reference)],
    )
    ws = wb.create_sheet(L.PLANNING.name)
    type_labels = {
        "LICENCE": "Licence",
        "CERTIFICAT": "Certificat",
        "MATERIEL": "Matériel",
        "APPLICATION": "Application",
        "CONTRAT": "Contrat",
    }
    plan = filter_renewals(renewal_items(assets, contracts), horizon_days=None)
    rows = []
    for it in plan:
        row = []
        for col in L.PLANNING.columns:
            v = it.get(col.key)
            if col.key == "type":
                v = type_labels.get(v, v)
            elif col.key == "status":
                v = lifecycle.STATUS_LABELS[v]
            row.append(v)
        rows.append(row)
    _write_table(ws, L.PLANNING.columns, rows)
    _write_dashboard(wb.create_sheet(L.DASHBOARD.name), await build_dashboard(db))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
