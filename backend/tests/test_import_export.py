import io
import sys
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
from generate_demo_excel import build  # noqa: E402

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def wb_bytes(sheets: dict[str, list[list]]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def upload(client, content: bytes, name="Suivi_Licence.xlsx"):
    return await client.post("/api/v1/import/excel", files={"file": (name, content, XLSX)})


LIC_HEAD = [
    "Référence",
    "Nom du logiciel",
    "Éditeur / Fournisseur",
    "Quantité totale",
    "Quantité utilisée",
    "Date d'expiration",
    "Criticité",
]


async def test_import_demo_file(admin, tmp_path):
    path = build(tmp_path / "Suivi_Licence.xlsx")
    r = await upload(admin, path.read_bytes())
    assert r.status_code == 200, r.text
    rep = r.json()
    assert rep["rows_imported"] > 40
    assert rep["missing_sheets"] == []
    assert rep["errors_count"] >= 1  # ligne sans référence ni nom
    msgs = " ".join(w["message"] for w in rep["warnings"])
    assert "doublon" in msgs and "date invalide" in msgs
    dash = (await admin.get("/api/v1/dashboard")).json()
    assert dash["kpis"]["licences"] == 16
    assert dash["licenses"]["used"] > 0

    # Ré-import = UPSERT, aucune création ni affectation en double
    used_before = dash["licenses"]["used"]
    rep2 = (await upload(admin, path.read_bytes())).json()
    assert rep2["created"] == 0 and rep2["updated"] == rep["rows_imported"]
    dash2 = (await admin.get("/api/v1/dashboard")).json()
    assert dash2["kpis"]["total_assets"] == dash["kpis"]["total_assets"]
    assert dash2["licenses"]["used"] == used_before


async def test_missing_sheet_is_reported(admin):
    content = wb_bytes(
        {"Licences_Logicielles": [LIC_HEAD, ["L1", "Office", "Microsoft", 10, 2, date.today() + timedelta(days=200), "Haute"]]}
    )
    rep = (await upload(admin, content)).json()
    assert rep["rows_imported"] == 1
    assert "Certificats" in rep["missing_sheets"] and "Contrats_Fournisseurs" in rep["missing_sheets"]


async def test_invalid_date_is_warning_not_blocking(admin):
    content = wb_bytes({"Licences_Logicielles": [LIC_HEAD, ["L1", "Office", "Microsoft", 10, 0, "32/13/2026", "Haute"]]})
    rep = (await upload(admin, content)).json()
    assert rep["rows_imported"] == 1
    assert any("date invalide" in w["message"] for w in rep["warnings"])
    asset = (await admin.get("/api/v1/assets")).json()["items"][0]
    assert asset["end_date"] is None and asset["status"] == "INCONNU"


async def test_duplicates_reported(admin):
    rows = [LIC_HEAD, ["L1", "Office", "Microsoft", 10, 0, None, None], ["L1", "Office bis", "Microsoft", 5, 0, None, None]]
    rep = (await upload(admin, wb_bytes({"Licences_Logicielles": rows}))).json()
    assert rep["rows_imported"] == 1 and rep["rows_skipped"] == 1
    assert any("doublon" in w["message"] for w in rep["warnings"])


async def test_no_business_sheet_rejected(admin):
    r = await upload(admin, wb_bytes({"Feuil1": [["a", "b"]]}))
    assert r.status_code == 422


async def test_non_excel_rejected(admin):
    r = await admin.post("/api/v1/import/excel", files={"file": ("data.csv", b"a,b", "text/csv")})
    assert r.status_code == 415


async def test_viewer_cannot_import(viewer):
    r = await upload(viewer, wb_bytes({"Licences_Logicielles": [LIC_HEAD]}))
    assert r.status_code == 403


async def test_export_roundtrip(admin, tmp_path):
    await upload(admin, build(tmp_path / "f.xlsx").read_bytes())
    r = await admin.get("/api/v1/export/excel")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == [
        "Licences_Logicielles",
        "Certificats",
        "Materiels",
        "Applications",
        "Contrats_Fournisseurs",
        "Planning_Renouvellement",
        "Dashboard",
    ]
    assert wb["Licences_Logicielles"].max_row == 17  # en-tête + 16 licences
    # Le fichier exporté est ré-importable sans création
    rep = (await upload(admin, r.content)).json()
    assert rep["created"] == 0 and rep["rows_imported"] > 40
