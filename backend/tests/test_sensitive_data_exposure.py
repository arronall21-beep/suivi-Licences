"""Correctifs de sécurité du lot RBAC : aucune donnée sensible ne doit sortir sans la permission correspondante.

Couvre les constats de l'audit : clés de licence (liste, détail, écriture, export), destinataires et erreurs SMTP des
alertes, historique des imports, documentation de l'API en production. Les tests portent sur le contenu réel des réponses,
pas seulement sur les codes HTTP.
"""

import io
import json
from datetime import date

from fastapi import FastAPI
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook

from app.core.config import Settings
from app.core.permissions import ROLE_PERMISSIONS, has_permission
from app.main import docs_urls
from app.models import ImportLog, Notification

API = "/api/v1"
KEY = "ABCD-1234-SECRET-KEY-9876"
RECIPIENT = "destinataire.secret@sbee.bj"
SMTP_ERROR = "535 5.7.8 authentification refusée pour smtp-user"


async def create_licence(admin, **extra) -> int:
    body = {"category": "LICENCE", "name": "Suite bureautique", "total_quantity": 5, "license_key": KEY, **extra}
    r = await admin.post(f"{API}/assets", json=body)
    assert r.status_code == 201, r.text
    return r.json()["id"]


def sheet_headers(content: bytes) -> dict[str, list[str]]:
    wb = load_workbook(io.BytesIO(content))
    return {ws.title: [c.value for c in ws[1]] for ws in wb.worksheets}


def all_cell_values(content: bytes) -> list:
    wb = load_workbook(io.BytesIO(content))
    return [c.value for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value is not None]


# --- permission ------------------------------------------------------------------------------------------------------


def test_secrets_permission_is_not_granted_by_read_permissions():
    assert not has_permission("VIEWER", "secrets:read")
    assert has_permission("MANAGER", "secrets:read") and has_permission("ADMIN", "secrets:read")
    # data:read et export:run seuls (VIEWER) ne suffisent jamais
    assert {"data:read", "export:run"} <= ROLE_PERMISSIONS["VIEWER"]


# --- clés de licence : lecture ---------------------------------------------------------------------------------------


async def test_viewer_never_receives_the_license_key(admin, viewer):
    asset_id = await create_licence(admin)

    detail = await viewer.get(f"{API}/assets/{asset_id}")
    assert detail.status_code == 200
    assert detail.json()["license_key"] is None
    assert detail.json()["has_license_key"] is True  # l'existence est indiquée, jamais la valeur

    listing = await viewer.get(f"{API}/assets", params={"size": 100})
    assert listing.status_code == 200
    assert [i["license_key"] for i in listing.json()["items"]] == [None]
    for r in (detail, listing):
        assert KEY not in r.text

    # les autres routes de lecture ne la contiennent pas non plus
    for path in ("/search?q=Suite", "/dashboard", "/renewals", "/assignments", "/alerts", "/notifications"):
        r = await viewer.get(API + path)
        assert r.status_code == 200 and KEY not in r.text, path


async def test_manager_and_admin_still_receive_the_license_key(admin, manager):
    asset_id = await create_licence(admin)
    for client in (admin, manager):
        assert (await client.get(f"{API}/assets/{asset_id}")).json()["license_key"] == KEY
        items = (await client.get(f"{API}/assets", params={"size": 100})).json()["items"]
        assert items[0]["license_key"] == KEY and items[0]["has_license_key"] is True


async def test_asset_without_key_reports_no_key(admin, viewer):
    asset_id = await create_licence(admin, license_key=None)
    body = (await viewer.get(f"{API}/assets/{asset_id}")).json()
    assert body["license_key"] is None and body["has_license_key"] is False


async def test_write_responses_do_not_leak_the_key_to_roles_without_the_permission(admin, monkeypatch):
    """Défense en profondeur : un rôle qui pourrait écrire sans `secrets:read` ne reçoit pas la clé en réponse."""
    monkeypatch.setitem(ROLE_PERMISSIONS, "MANAGER", frozenset(ROLE_PERMISSIONS["MANAGER"] - {"secrets:read"}))
    from tests.conftest import make_client

    manager = await make_client("manager@example.com", "manager123")
    try:
        created = await manager.post(f"{API}/assets", json={"category": "LICENCE", "name": "Outil", "license_key": KEY})
        assert created.status_code == 201 and created.json()["license_key"] is None and KEY not in created.text
        asset_id = created.json()["id"]
        updated = await manager.put(f"{API}/assets/{asset_id}", json={"category": "LICENCE", "name": "Outil v2"})
        assert updated.status_code == 200 and KEY not in updated.text
        # la modification sans permission n'efface pas la clé existante
        assert (await admin.get(f"{API}/assets/{asset_id}")).json()["license_key"] == KEY
    finally:
        await manager.aclose()


async def test_update_by_authorized_role_can_change_and_clear_the_key(admin, manager):
    asset_id = await create_licence(admin)
    body = {"category": "LICENCE", "name": "Suite bureautique", "license_key": "NOUVELLE-CLE"}
    assert (await manager.put(f"{API}/assets/{asset_id}", json=body)).json()["license_key"] == "NOUVELLE-CLE"
    body["license_key"] = None
    assert (await manager.put(f"{API}/assets/{asset_id}", json=body)).json()["license_key"] is None


async def test_audit_log_never_stores_the_key_value(admin):
    asset_id = await create_licence(admin)
    await admin.put(
        f"{API}/assets/{asset_id}", json={"category": "LICENCE", "name": "Suite bureautique", "license_key": "AUTRE-CLE-777"}
    )
    r = await admin.get(f"{API}/admin/audit", params={"size": 200})
    assert KEY not in r.text and "AUTRE-CLE-777" not in r.text


# --- clés de licence : exports --------------------------------------------------------------------------------------


async def test_exports_omit_the_key_column_without_the_permission(admin, viewer):
    await create_licence(admin)
    for path in ("excel", "licences", "contrats", "echeances", "planning"):
        r = await viewer.get(f"{API}/export/{path}")
        assert r.status_code == 200, path
        assert KEY not in all_cell_values(r.content), path
        assert "Clé de licence" not in {h for headers in sheet_headers(r.content).values() for h in headers}, path


async def test_exports_keep_the_key_for_authorized_roles(admin, manager):
    await create_licence(admin)
    for client in (admin, manager):
        for path in ("excel", "licences"):
            r = await client.get(f"{API}/export/{path}")
            assert KEY in all_cell_values(r.content), path
            assert "Clé de licence" in {h for headers in sheet_headers(r.content).values() for h in headers}, path


async def test_export_is_audited_with_key_inclusion_flag(admin, viewer):
    await create_licence(admin)
    await admin.get(f"{API}/export/licences")
    await viewer.get(f"{API}/export/licences")
    items = (await admin.get(f"{API}/admin/audit", params={"action": "EXPORT_EXCEL", "size": 10})).json()["items"]
    flags = sorted(json.loads(i["details"])["cles_incluses"] for i in items)
    assert flags == [False, True]


# --- alertes et notifications d'envoi ------------------------------------------------------------------------------


async def seed_alert_log(db):
    db.add(
        Notification(
            target_type="ASSET",
            target_id=1,
            target_reference="LIC-001",
            target_name="Suite bureautique",
            threshold=30,
            due_date=date(2030, 1, 1),
            days_remaining=30,
            recipients=RECIPIENT,
            delivery_status="ERREUR",
            error=SMTP_ERROR,
        )
    )
    await db.commit()


async def test_alert_delivery_log_hides_recipients_and_errors_from_read_only_roles(db, admin, manager, viewer):
    await seed_alert_log(db)
    for client in (viewer, manager):
        r = await client.get(f"{API}/notifications")
        assert r.status_code == 200 and len(r.json()) == 1
        row = r.json()[0]
        assert row["recipients"] is None and row["error"] is None
        assert row["delivery_status"] == "ERREUR" and row["target_reference"] == "LIC-001"  # le suivi métier reste visible
        assert RECIPIENT not in r.text and SMTP_ERROR not in r.text
    full = (await admin.get(f"{API}/notifications")).json()[0]
    assert full["recipients"] == RECIPIENT and full["error"] == SMTP_ERROR


async def test_alerts_overview_hides_configured_recipients_from_read_only_roles(admin, manager, viewer):
    cfg = (await admin.get(f"{API}/admin/settings/alerts")).json()
    cfg["recipients"]["custom"] = [RECIPIENT]
    assert (await admin.put(f"{API}/admin/settings/alerts", json=cfg)).status_code == 200

    for client in (viewer, manager):
        r = await client.get(f"{API}/alerts")
        assert r.status_code == 200
        assert r.json()["recipients"] == [] and r.json()["recipients_visible"] is False
        assert RECIPIENT not in r.text
        assert "items" in r.json() and "thresholds" in r.json()  # le tableau des échéances reste disponible
    full = (await admin.get(f"{API}/alerts")).json()
    assert full["recipients_visible"] is True and RECIPIENT in full["recipients"]


# --- historique des imports -----------------------------------------------------------------------------------------


async def test_import_history_is_reserved_to_importers(db, admin, manager, viewer):
    report = {"errors": [f"Ligne 4 : clé « {KEY} » invalide"], "warnings": []}
    log = ImportLog(
        filename="export_confidentiel.xlsx",
        rows_analyzed=1,
        rows_imported=0,
        rows_skipped=1,
        warnings_count=0,
        errors_count=1,
        report=json.dumps(report),
        imported_by="admin@example.com",
        file_size=10,
        status="ERRORS",
    )
    db.add(log)
    await db.commit()
    await db.refresh(log)

    for client in (viewer, manager):
        for path in ("/import/history", f"/import/history/{log.id}"):
            r = await client.get(API + path)
            assert r.status_code == 403, path
            assert "export_confidentiel" not in r.text and KEY not in r.text
    listing = await admin.get(f"{API}/import/history")
    assert listing.status_code == 200 and listing.json()[0]["filename"] == "export_confidentiel.xlsx"
    assert KEY in (await admin.get(f"{API}/import/history/{log.id}")).text


# --- documentation de l'API ----------------------------------------------------------------------------------------


def test_docs_are_enabled_only_in_development_like_environments():
    enabled = {env: Settings(environment=env).api_docs_enabled for env in ("development", "test", "local", "dev")}
    assert all(enabled.values()), enabled
    for env in ("production", "Production", " production ", "prod", "staging", ""):
        assert Settings(environment=env).api_docs_enabled is False, env  # refus par défaut pour tout autre environnement


async def test_disabled_docs_return_404_on_every_documentation_url():
    app = FastAPI(**docs_urls(False))
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as c:
        for path in ("/api/docs", "/api/openapi.json", "/redoc", "/docs/oauth2-redirect", "/docs", "/openapi.json"):
            assert (await c.get(path)).status_code == 404, path


async def test_enabled_docs_are_served_in_development():
    app = FastAPI(**docs_urls(True))
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as c:
        assert (await c.get("/api/docs")).status_code == 200
        assert (await c.get("/api/openapi.json")).status_code == 200


def test_application_uses_the_settings_driven_docs_configuration():
    from app.main import app as application
    from app.main import settings

    expected = docs_urls(settings.api_docs_enabled)
    assert application.docs_url == expected["docs_url"]
    assert application.openapi_url == expected["openapi_url"]
    assert application.redoc_url == expected["redoc_url"]
