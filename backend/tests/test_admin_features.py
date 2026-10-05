import io
import json
from datetime import date, timedelta

from openpyxl import Workbook, load_workbook

from app.services import backup

API = "/api/v1"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def iso(days: int) -> str:
    return (date.today() + timedelta(days=days)).isoformat()


def workbook(rows, sheet="Licences_Logicielles"):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(["ID", "Logiciel", "Nb licences", "Date expiration"])
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def audit_entries(admin, **params):
    return (await admin.get(f"{API}/admin/audit", params={"size": 200, **params})).json()


# ---------- Audit ----------
async def test_audit_records_business_and_admin_actions(admin, manager):
    a = (await admin.post(f"{API}/assets", json={"category": "LICENCE", "name": "Office", "total_quantity": 10})).json()
    await manager.put(f"{API}/assets/{a['id']}", json={"category": "LICENCE", "name": "Office 365", "total_quantity": 10})
    asg = (await manager.post(f"{API}/assignments", json={"asset_id": a["id"], "assigned_to_user": "awa", "quantity": 3})).json()
    refused = await manager.post(f"{API}/assignments", json={"asset_id": a["id"], "assigned_to_user": "x", "quantity": 99})
    assert refused.status_code == 422
    await manager.delete(f"{API}/assignments/{asg['id']}")
    await admin.delete(f"{API}/assets/{a['id']}")
    await admin.put(
        f"{API}/admin/settings/general",
        json={
            "org_name": "SBEE",
            "app_name": "Patrimoine",
            "currency": "XOF",
            "timezone": "UTC",
            "admin_email": None,
            "critical_days": 30,
            "alert_days": 90,
        },
    )

    entries = (await audit_entries(admin))["items"]
    actions = [e["action"] for e in entries]
    for expected in (
        "ASSET_CREATE",
        "ASSET_UPDATE",
        "ASSIGNMENT_CREATE",
        "ASSIGNMENT_DELETE",
        "ASSET_DELETE",
        "SETTINGS_UPDATE",
        "LOGIN_SUCCESS",
    ):
        assert expected in actions, expected
    failure = next(e for e in entries if e["action"] == "ASSIGNMENT_CREATE" and e["result"] == "FAILURE")
    assert "Sur-allocation" in failure["details"] and failure["user_email"] == "manager@example.com"
    create = next(e for e in entries if e["action"] == "ASSET_CREATE")
    assert create["user_email"] == "admin@example.com" and create["entity_label"].startswith("LIC-001") and create["ip_address"]
    update = next(e for e in entries if e["action"] == "ASSET_UPDATE")
    assert json.loads(update["details"])["champs"] == ["name"] and update["user_email"] == "manager@example.com"
    unassign = next(e for e in entries if e["action"] == "ASSIGNMENT_DELETE")
    assert json.loads(unassign["details"])["cible"] == "awa"


async def test_audit_user_admin_actions_and_no_secrets(admin):
    pwd = "Secret123"
    u = (
        await admin.post(
            f"{API}/users", json={"first_name": "A", "last_name": "B", "email": "ab@sbee.bj", "role": "VIEWER", "password": pwd}
        )
    ).json()
    body = {"first_name": "A", "last_name": "B", "email": "ab@sbee.bj", "phone": None, "department": "DSI", "role": "MANAGER"}
    await admin.put(f"{API}/users/{u['id']}", json=body)
    await admin.post(f"{API}/users/{u['id']}/deactivate")
    await admin.post(f"{API}/users/{u['id']}/reactivate")
    await admin.post(f"{API}/users/{u['id']}/reset-password", json={"password": "Nouveau456"})
    await admin.put(
        f"{API}/admin/settings/smtp",
        json={
            "host": "smtp.sbee.bj",
            "port": 587,
            "security": "STARTTLS",
            "username": "u",
            "password": "TopSecretSMTP1",
            "from_email": "a@sbee.bj",
        },
    )
    await admin.post(f"{API}/auth/login", json={"email": "ab@sbee.bj", "password": "mauvais"})

    raw = json.dumps((await audit_entries(admin))["items"])
    for secret in (pwd, "Nouveau456", "TopSecretSMTP1"):
        assert secret not in raw
    actions = {e["action"] for e in (await audit_entries(admin))["items"]}
    assert {
        "USER_CREATE",
        "USER_UPDATE",
        "USER_ROLE_CHANGE",
        "USER_DEACTIVATE",
        "USER_REACTIVATE",
        "USER_PASSWORD_RESET",
        "SMTP_UPDATE",
        "LOGIN_FAILURE",
    } <= actions
    role = next(e for e in (await audit_entries(admin))["items"] if e["action"] == "USER_ROLE_CHANGE")
    assert json.loads(role["details"]) == {"de": "VIEWER", "vers": "MANAGER"}
    smtp = next(e for e in (await audit_entries(admin))["items"] if e["action"] == "SMTP_UPDATE")
    assert "password (remplacé)" in json.loads(smtp["details"])["champs"]


async def test_audit_filters_and_pagination(admin, manager):
    for i in range(3):
        await manager.post(f"{API}/assets", json={"category": "LICENCE", "name": f"L{i}"})
    await admin.post(f"{API}/vendors", json={"name": "V"})
    assert (await audit_entries(admin, action="ASSET_CREATE"))["total"] == 3
    assert (await audit_entries(admin, user="manager"))["total"] >= 3
    assert (await audit_entries(admin, entity_type="VENDOR"))["total"] == 1
    assert (await audit_entries(admin, q="LIC-002"))["total"] == 1
    assert (await audit_entries(admin, result="FAILURE"))["total"] == 0
    today = date.today().isoformat()
    assert (await audit_entries(admin, date_from=today, date_to=today))["total"] > 0
    assert (await audit_entries(admin, date_to=iso(-2)))["total"] == 0
    page = (await admin.get(f"{API}/admin/audit", params={"size": 2, "page": 1})).json()
    assert len(page["items"]) == 2 and page["total"] > 2
    assert len((await admin.get(f"{API}/admin/audit/actions")).json()) > 10


async def test_audit_is_admin_only(manager, viewer):
    for client in (manager, viewer):
        assert (await client.get(f"{API}/admin/audit")).status_code == 403
        assert (await client.get(f"{API}/admin/audit/actions")).status_code == 403
        assert (await client.get(f"{API}/admin/system")).status_code == 403


# ---------- Historique d'import + notifications ----------
async def test_import_history_and_inbox(admin, manager):
    content = workbook([["LIC-001", "Office", 5, None], [None, None, 9, None]])
    r = await admin.post(f"{API}/import/excel", files={"file": ("suivi.xlsx", content, XLSX)})
    assert r.status_code == 200
    bad = await admin.post(f"{API}/import/excel", files={"file": ("rien.xlsx", workbook([], sheet="Feuil1"), XLSX)})
    assert bad.status_code == 422
    csv = await admin.post(f"{API}/import/excel", files={"file": ("a.csv", b"a,b", "text/csv")})
    assert csv.status_code == 415

    history = (await admin.get(f"{API}/import/history")).json()
    assert [h["status"] for h in history] == ["FAILED", "FAILED", "ERRORS"]
    ok = history[2]
    assert ok["filename"] == "suivi.xlsx" and ok["file_size"] == len(content) and ok["rows_imported"] == 1
    assert ok["imported_by"] == "admin@example.com" and ok["errors_count"] == 1
    failed_report = (await admin.get(f"{API}/import/history/{history[0]['id']}")).json()
    assert failed_report["errors"][0]["message"].startswith("Format attendu")
    assert (await manager.get(f"{API}/import/history")).status_code == 200  # consultation ouverte à tous les connectés

    inbox = (await admin.get(f"{API}/inbox")).json()
    kinds = [n["kind"] for n in inbox["items"]]
    assert kinds.count("IMPORT_DONE") == 1 and kinds.count("IMPORT_ERROR") == 2
    assert next(n for n in inbox["items"] if n["kind"] == "IMPORT_ERROR")["priority"] == "HIGH"
    # destinées aux administrateurs : invisibles pour un gestionnaire
    assert (await manager.get(f"{API}/inbox")).json()["total"] == 0

    audit = [e for e in (await audit_entries(admin))["items"] if e["action"] == "IMPORT_EXCEL"]
    assert sorted(e["result"] for e in audit) == ["FAILURE", "FAILURE", "SUCCESS"]


async def test_inbox_read_state_is_per_user(admin, manager, viewer):
    await admin.post(
        f"{API}/assets", json={"category": "CONTRAT" if False else "LICENCE", "name": "Expirée", "end_date": iso(-1)}
    )
    await admin.post(f"{API}/alerts/run")
    inbox_admin = (await admin.get(f"{API}/inbox")).json()
    assert inbox_admin["unread"] == 1 and inbox_admin["items"][0]["read"] is False
    nid = inbox_admin["items"][0]["id"]
    assert (await manager.get(f"{API}/inbox/unread-count")).json()["unread"] == 1
    assert (await manager.post(f"{API}/inbox/{nid}/read")).json()["unread"] == 0
    assert (await manager.get(f"{API}/inbox", params={"unread_only": True})).json()["total"] == 0
    assert (await viewer.get(f"{API}/inbox/unread-count")).json()["unread"] == 1  # l'état de lecture est personnel
    assert (await admin.get(f"{API}/inbox", params={"unread_only": True})).json()["total"] == 1
    assert (await viewer.post(f"{API}/inbox/read-all")).json() == {"unread": 0}
    assert (await admin.post(f"{API}/inbox/999/read")).status_code == 404
    assert (await admin.get(f"{API}/inbox", params={"priority": "LOW"})).json()["total"] == 0


async def test_inbox_alert_types(admin):
    await admin.post(f"{API}/contracts", json={"scope": "Maintenance", "end_date": iso(20)})
    await admin.post(f"{API}/assets", json={"category": "LICENCE", "name": "Lic critique", "end_date": iso(20)})
    await admin.post(f"{API}/assets", json={"category": "CERTIFICAT", "name": "cert", "end_date": iso(-5)})
    await admin.post(f"{API}/alerts/run")
    kinds = {n["kind"] for n in (await admin.get(f"{API}/inbox")).json()["items"]}
    assert {"CONTRACT_EXPIRING", "LICENSE_CRITICAL", "CERTIFICATE_EXPIRED"} <= kinds
    run2 = (await admin.post(f"{API}/alerts/run")).json()
    assert run2["new"] == 0 and (await admin.get(f"{API}/inbox")).json()["total"] == 3  # idempotent


# ---------- Exports ----------
async def test_scoped_exports(admin, viewer):
    await admin.post(f"{API}/assets", json={"category": "LICENCE", "name": "Office", "end_date": iso(10), "annual_cost": 1000})
    await admin.post(f"{API}/assets", json={"category": "CERTIFICAT", "name": "vpn", "end_date": iso(400)})
    await admin.post(f"{API}/contracts", json={"scope": "Support", "end_date": iso(60)})
    expected = {
        "licences": ["Licences_Logicielles"],
        "contrats": ["Contrats_Fournisseurs"],
        "echeances": ["Echeances"],
        "planning": ["Planning_Renouvellement"],
    }
    for scope, sheets in expected.items():
        r = await viewer.get(f"{API}/export/{scope}")
        assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
        wb = load_workbook(io.BytesIO(r.content))
        assert wb.sheetnames == sheets, scope
    due = load_workbook(io.BytesIO((await admin.get(f"{API}/export/echeances")).content))["Echeances"]
    assert due.max_row == 3  # en-tête + licence (10 j) + contrat (60 j) ; le certificat à 400 j est exclu
    plan = load_workbook(io.BytesIO((await admin.get(f"{API}/export/planning", params={"horizon_days": 90})).content))
    assert plan["Planning_Renouvellement"].max_row == 3
    full = load_workbook(io.BytesIO((await admin.get(f"{API}/export/excel")).content))
    assert len(full.sheetnames) == 7
    lic = load_workbook(io.BytesIO((await admin.get(f"{API}/export/licences")).content))["Licences_Logicielles"]
    header = [c.value for c in lic[1]]
    assert (
        header[:2] == ["Référence", "Nom du logiciel"]
        and "FCFA" in lic.cell(row=2, column=header.index("Coût annuel") + 1).number_format
    )
    actions = [e["details"] for e in (await audit_entries(admin))["items"] if e["action"] == "EXPORT_EXCEL"]
    assert any("licences" in d for d in actions)


# ---------- Dashboard : top des risques ----------
async def test_dashboard_top_risks(admin):
    mk = lambda **kw: admin.post(f"{API}/assets", json={"category": "LICENCE", "total_quantity": 100, **kw})  # noqa: E731
    await mk(name="Basse expirée", criticality="Basse", end_date=iso(-10))
    await mk(name="Critique expirée", criticality="Critique", end_date=iso(-10))
    await mk(name="Haute 20j", criticality="Haute", end_date=iso(20))
    await mk(name="OK loin", criticality="Critique", end_date=iso(300))
    await mk(name="Sans date", criticality="Critique")
    big = (await mk(name="Critique 20j très utilisée", criticality="Critique", end_date=iso(20))).json()
    await admin.post(f"{API}/assignments", json={"asset_id": big["id"], "assigned_to_department": "DSI", "quantity": 80})
    risks = (await admin.get(f"{API}/dashboard")).json()["top_risks"]
    assert [r["name"] for r in risks] == ["Critique expirée", "Critique 20j très utilisée", "Haute 20j", "Basse expirée"]
    assert [r["score"] for r in risks] == [20, 18, 12, 5]  # 4×5 ; 4×4+2 ; 3×4 ; 1×5
    assert risks[1]["reasons"] == ["criticité Critique (×4)", "échéance dans 20 j (×4)", "utilisation 80 % (+2)"]


# ---------- Recherche globale ----------
async def test_global_search_groups_and_permissions(admin, viewer):
    v = (await admin.post(f"{API}/vendors", json={"name": "Cisco Systems", "contact_person": "Marc"})).json()
    await admin.post(f"{API}/assets", json={"category": "MATERIEL", "name": "Switch cœur", "vendor_id": v["id"]})
    await admin.post(f"{API}/assets", json={"category": "LICENCE", "name": "Figma"})
    await admin.post(f"{API}/contracts", json={"scope": "SmartNet Cisco", "vendor_id": v["id"]})
    r = (await admin.get(f"{API}/search", params={"q": "cisco"})).json()
    assert [a["label"] for a in r["assets"]] == ["MAT-001 — Switch cœur"]
    assert r["contracts"][0]["label"].startswith("CTR-001") and r["vendors"][0]["label"] == "VEN-001 — Cisco Systems"
    assert r["users"] == []
    assert [a["label"] for a in (await admin.get(f"{API}/search", params={"q": "LIC-001"})).json()["assets"]] == [
        "LIC-001 — Figma"
    ]
    assert (await admin.get(f"{API}/search", params={"q": "admin@"})).json()["users"][0]["link"].startswith("/admin/utilisateurs")
    assert (await viewer.get(f"{API}/search", params={"q": "admin@"})).json()["users"] == []  # utilisateurs : ADMIN seulement
    assert (await admin.get(f"{API}/search", params={"q": "a"})).status_code == 422


# ---------- Système ----------
async def test_system_info_and_backup(admin, tmp_path, monkeypatch):
    monkeypatch.setattr(backup.settings, "backup_dir", str(tmp_path / "absent"))
    s = (await admin.get(f"{API}/admin/system")).json()
    assert s["version"] == "1.1.0" and s["database"]["status"] == "ok" and s["smtp"]["configured"] is False
    assert s["users"] == {"active": 3, "inactive": 0} and s["backup"]["status"] == "NON_CONFIGURE"
    assert s["last_import"] is None and s["last_alert_sent"] is None

    monkeypatch.setattr(backup.settings, "backup_dir", str(tmp_path))
    assert (await admin.get(f"{API}/admin/system")).json()["backup"]["status"] == "AUCUN"
    dump = tmp_path / "suivi_20261005.sql"
    dump.write_text("-- dump")
    b = (await admin.get(f"{API}/admin/system")).json()["backup"]
    assert b["status"] == "OK" and b["filename"] == "suivi_20261005.sql" and b["count"] == 1
    old = dump.stat().st_mtime - 5 * 24 * 3600
    import os

    os.utime(dump, (old, old))
    assert (await admin.get(f"{API}/admin/system")).json()["backup"]["status"] == "ANCIEN"
